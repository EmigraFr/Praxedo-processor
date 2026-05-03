#!/usr/bin/env python3
"""
Praxedo Queries — Couche de requêtes réutilisable
==================================================
Fournit une API simple pour interroger la base :
  - search_interventions(filters) : recherche multi-critères
  - get_intervention(id)          : détail d'une intervention
  - get_history(refint)           : historique des versions
  - get_distinct_values(column)   : valeurs distinctes (pour listes déroulantes)
  - global_stats()                : KPI pour le dashboard

Toutes les fonctions utilisent des requêtes paramétrées (pas d'injection SQL).
"""

from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


# =============================================================================
# Colonnes affichables (clé technique -> libellé lisible)
# =============================================================================

DISPLAY_COLUMNS = [
    ("refint",             "Refint"),
    ("nd",                 "ND"),
    ("poi",                "POI"),
    ("tipologie",          "Tipologie"),
    ("arc",                "ARC"),
    ("agence",             "Agence"),
    ("technicien_nom",     "Technicien"),
    ("client",             "Client"),
    ("date_planifiee",     "Planifiée"),
    ("caff_nom",           "CAFF"),
    ("rai_nom",            "RAI"),
    ("email_rpe",          "Email RPE"),
    ("nb_versions",        "Vers."),
    ("updated_at",         "Maj"),
]

FILTER_FIELDS = [
    # (clé, libellé, type, colonne SQL)
    ("nd",              "ND",                 "text_like", "nd"),
    ("poi",             "POI",                "text_like", "poi"),
    ("refint",          "Refint",             "text_like", "refint"),
    ("tipologie",       "Tipologie",          "exact",     "tipologie"),
    ("arc",             "ARC",                "bool",      "arc"),
    ("agence",          "Agence",             "exact",     "agence"),
    ("technicien_nom",  "Technicien",         "text_like", "technicien_nom"),
    ("client",          "Client",             "text_like", "client"),
    ("caff_nom",        "CAFF",               "text_like", "caff_nom"),
    ("rai_nom",         "RAI",                "text_like", "rai_nom"),
    ("email_rpe",       "Email RPE",          "text_like", "email_rpe"),
    ("date_from",       "Planifiée après",    "date_ge",   "date_planifiee"),
    ("date_to",         "Planifiée avant",    "date_le",   "date_planifiee"),
]

SEARCH_GLOBAL_COLUMNS = [
    "refint", "nd", "poi", "description", "caff_nom", "caff_tel",
    "rai_nom", "rai_tel", "email_rpe", "telephones", "client",
    "technicien_nom", "technicien_prenom",
]


# =============================================================================
# Classe principale
# =============================================================================

class PraxedoQueries:
    """Wrapper de requêtes sur la base Praxedo (lecture seule)."""

    def __init__(self, db_path: Union[str, Path]):
        self.db_path = Path(db_path)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row

    def close(self):
        self.conn.close()

    def __enter__(self): return self
    def __exit__(self, *a): self.close()

    # ------------------------------------------------------------------
    # RECHERCHE MULTI-CRITÈRES
    # ------------------------------------------------------------------
    def search_interventions(self, filters: Dict[str, Any] | None = None,
                             global_search: Optional[str] = None,
                             order_by: str = "updated_at",
                             order_dir: str = "DESC",
                             limit: int = 500,
                             offset: int = 0) -> Tuple[List[sqlite3.Row], int]:
        """
        Recherche avec filtres. Retourne (résultats, nb_total).
        - `filters` : dict des filtres (cf. FILTER_FIELDS).
        - `global_search` : texte libre cherché dans plusieurs colonnes.
        - `order_by` : colonne SQL de tri.
        - `order_dir` : ASC ou DESC.
        """
        filters = filters or {}
        where_parts = []
        params: List[Any] = []

        # Filtres champ par champ
        for key, label, ftype, sql_col in FILTER_FIELDS:
            val = filters.get(key)
            if val in (None, "", "(tous)"):
                continue
            if ftype == "text_like":
                where_parts.append(f"LOWER({sql_col}) LIKE ?")
                params.append(f"%{str(val).lower()}%")
            elif ftype == "exact":
                where_parts.append(f"{sql_col} = ?")
                params.append(str(val))
            elif ftype == "bool":
                # 1 pour oui, 0 pour non
                where_parts.append(f"{sql_col} = ?")
                params.append(1 if str(val).lower() in ("1", "oui", "yes", "true") else 0)
            elif ftype == "date_ge":
                where_parts.append(f"{sql_col} >= ?")
                params.append(str(val))
            elif ftype == "date_le":
                where_parts.append(f"{sql_col} <= ?")
                params.append(str(val))

        # Recherche globale en texte libre
        if global_search and str(global_search).strip():
            like = f"%{str(global_search).strip().lower()}%"
            parts = [f"LOWER(COALESCE({c},'')) LIKE ?" for c in SEARCH_GLOBAL_COLUMNS]
            where_parts.append("(" + " OR ".join(parts) + ")")
            params.extend([like] * len(SEARCH_GLOBAL_COLUMNS))

        where_sql = (" WHERE " + " AND ".join(where_parts)) if where_parts else ""

        # Sécurisation du tri (whitelist)
        allowed_order = {c for c, _ in DISPLAY_COLUMNS} | {
            "created_at", "updated_at", "nb_versions",
        }
        order_col = order_by if order_by in allowed_order else "updated_at"
        order_sql = f"ORDER BY {order_col} {'ASC' if order_dir == 'ASC' else 'DESC'}"

        # Comptage
        count_q = f"SELECT COUNT(*) AS n FROM interventions {where_sql}"
        total = self.conn.execute(count_q, params).fetchone()["n"]

        # Résultats paginés
        data_q = (f"SELECT * FROM interventions {where_sql} "
                  f"{order_sql} LIMIT ? OFFSET ?")
        rows = self.conn.execute(data_q, params + [limit, offset]).fetchall()

        return list(rows), total

    # ------------------------------------------------------------------
    # DÉTAIL + HISTORIQUE
    # ------------------------------------------------------------------
    def get_intervention(self, intervention_id: int) -> Optional[sqlite3.Row]:
        return self.conn.execute(
            "SELECT * FROM interventions WHERE id = ?", (intervention_id,)
        ).fetchone()

    def get_intervention_by_refint(self, refint: str) -> Optional[sqlite3.Row]:
        return self.conn.execute(
            "SELECT * FROM interventions WHERE refint = ?", (refint,)
        ).fetchone()

    def get_history(self, refint: str) -> List[sqlite3.Row]:
        """Toutes les versions d'un Refint, de la plus ancienne à la plus récente."""
        return list(self.conn.execute(
            "SELECT h.*, i.date_import AS import_date, f.nom_fichier AS source_file "
            "FROM interventions_historique h "
            "LEFT JOIN imports i ON i.id = h.import_id "
            "LEFT JOIN fichiers_sources f ON f.id = i.fichier_source_id "
            "WHERE h.refint = ? "
            "ORDER BY h.numero_version ASC",
            (refint,),
        ).fetchall())

    # ------------------------------------------------------------------
    # VALEURS DISTINCTES (pour listes déroulantes)
    # ------------------------------------------------------------------
    def get_distinct_values(self, column: str) -> List[str]:
        allowed = {"agence", "tipologie", "technicien_nom",
                   "statut", "type_inter", "act_prod"}
        if column not in allowed:
            return []
        rows = self.conn.execute(
            f"SELECT DISTINCT {column} AS v FROM interventions "
            f"WHERE {column} IS NOT NULL AND {column} != '' "
            f"ORDER BY {column}"
        ).fetchall()
        return [r["v"] for r in rows]

    # ------------------------------------------------------------------
    # STATISTIQUES GLOBALES (pour Phase 3)
    # ------------------------------------------------------------------
    def global_stats(self) -> dict:
        c = self.conn

        def scalar(q, *args):
            r = c.execute(q, args).fetchone()
            return list(r)[0] if r else 0

        return {
            "total":         scalar("SELECT COUNT(*) FROM interventions"),
            "std":           scalar("SELECT COALESCE(SUM(standard),0) FROM interventions"),
            "nstd":          scalar("SELECT COALESCE(SUM(non_standard),0) FROM interventions"),
            "arc":           scalar("SELECT COALESCE(SUM(arc),0) FROM interventions"),
            "agences":       scalar("SELECT COUNT(DISTINCT agence) FROM interventions WHERE agence IS NOT NULL"),
            "techniciens":   scalar("SELECT COUNT(DISTINCT technicien_nom) FROM interventions WHERE technicien_nom IS NOT NULL"),
            "clients":       scalar("SELECT COUNT(DISTINCT client) FROM interventions WHERE client IS NOT NULL"),
            "versions":      scalar("SELECT COUNT(*) FROM interventions_historique"),
            "doublons":      scalar("SELECT COUNT(*) FROM interventions WHERE nb_versions > 1"),
        }

    def stats_by_agence(self) -> List[sqlite3.Row]:
        return list(self.conn.execute(
            "SELECT * FROM v_stats_agence"
        ).fetchall())

    def stats_by_month(self, months: int = 12) -> List[sqlite3.Row]:
        return list(self.conn.execute(
            "SELECT * FROM v_stats_mensuelles "
            "ORDER BY mois DESC LIMIT ?", (months,)
        ).fetchall())

    def top_techniciens(self, n: int = 10) -> List[sqlite3.Row]:
        return list(self.conn.execute(
            "SELECT * FROM v_top_techniciens LIMIT ?", (n,)
        ).fetchall())


# =============================================================================
# EXPORT EXCEL DE RÉSULTATS
# =============================================================================

def export_rows_to_xlsx(rows: List[sqlite3.Row], output_path: Path,
                        title: str = "Résultats de recherche") -> Path:
    """Exporte une liste de résultats SQL vers un fichier Excel mis en forme."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultats"

    if not rows:
        ws["A1"] = "Aucun résultat"
        wb.save(output_path)
        return output_path

    # Utiliser les clés de la 1re ligne comme en-têtes
    keys = list(rows[0].keys())

    # En-têtes
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    center = Alignment(horizontal="center", vertical="center")
    for c, k in enumerate(keys, start=1):
        cell = ws.cell(row=1, column=c, value=k)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # Données
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, k in enumerate(keys, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row[k])

    # Largeurs de colonnes
    for c, k in enumerate(keys, start=1):
        letter = ws.cell(row=1, column=c).column_letter
        ws.column_dimensions[letter].width = max(10, min(40, len(k) + 4))

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path
